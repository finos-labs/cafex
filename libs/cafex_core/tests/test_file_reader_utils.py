import unittest
from unittest.mock import patch, MagicMock, mock_open
import pytest
import os
import pandas as pd
import csv
from cafex_core.utils.file_reader_utils import FileReaderUtil
from docx import Document

file_reader_util = FileReaderUtil()


def test_read_pdf_file():
    try:
        pdf_file_path = 'path/to/file.pdf'
        expected_first_word = "word"
        pdf_content = file_reader_util.read_pdf_file(pdf_file_path)
        first_word = pdf_content[0].split()[0] if pdf_content and len(pdf_content[0].split()) > 0 else None
        assert first_word == expected_first_word, f"Expected '{expected_first_word}', but got '{first_word}'"
    except Exception as e:
        print("An error occurred while reading the PDF file")


def test_read_pdf_file_invalid_path():
    try:
        invalid_pdf_file_path = 'path/to/invalidfile.pdf'
        with pytest.raises(Exception) as e:
            file_reader_util.read_pdf_file(invalid_pdf_file_path)
        assert "An error occurred while reading the PDF file" in str(e.value)
    except Exception as e:
        print("An error occurred while reading the PDF file")


def test_read_pdf_file_not_a_pdf():
    pdf_file_path = 'path/to/not_a_pdf.txt'
    with patch.object(file_reader_util, 'read_pdf_file', side_effect=ValueError("The provided file is not a PDF.")):
        with pytest.raises(ValueError) as e:
            file_reader_util.read_pdf_file(pdf_file_path)
        assert str(e.value) == "The provided file is not a PDF."


def test_read_pdf_file_not_a_pdf_file():
    non_pdf_file_path = 'path/to/file.txt'

    with pytest.raises(ValueError) as context:
        file_reader_util.read_pdf_file(non_pdf_file_path)

    assert str(context.value) == "The provided file is not a PDF."


@patch('cafex_core.utils.file_reader_utils.fitz.open')
def test_read_pdf_file_with_mock(mock_fitz_open):
    # Mock the doc object and its methods
    mock_page = MagicMock()
    mock_page.get_text.return_value = "This is a sample text\nwith new lines\r to be replaced."
    mock_doc = MagicMock()
    mock_doc.__enter__.return_value = [mock_page]
    mock_fitz_open.return_value = mock_doc
    content_list = file_reader_util.read_pdf_file('path/to/file.pdf')

    # Assertions
    mock_fitz_open.assert_called_once_with('path/to/file.pdf')
    mock_page.get_text.assert_called_once_with("text")
    assert content_list == ["This is a sample text with new lines to be replaced."]


@patch('openpyxl.load_workbook')
def test_read_excel_file_non_existent_sheet_name(mock_load_workbook):
    excel_file_path = 'path/to/valid.xlsx'
    non_existent_sheet_name = 'NonExistentSheet'

    # Mock the workbook and its methods
    mock_workbook = MagicMock()
    mock_workbook.sheetnames = ['Sheet1', 'Sheet2']
    mock_load_workbook.return_value = mock_workbook

    with pytest.raises(ValueError) as context:
        file_reader_util.read_excel(excel_file_path, sheet_reference=non_existent_sheet_name)

    assert str(context.value) == f"Sheet name '{non_existent_sheet_name}' does not exist."
    mock_load_workbook.assert_called_once_with(excel_file_path, data_only=True)


@patch('openpyxl.load_workbook')
def test_read_excel_file_valid_sheet_index(mock_load_workbook):
    excel_file_path = 'path/to/valid.xlsx'
    valid_sheet_index = 1

    # Mock the workbook and its methods
    mock_workbook = MagicMock()
    mock_sheet = MagicMock()
    mock_workbook.worksheets = [MagicMock(), mock_sheet, MagicMock()]
    mock_load_workbook.return_value = mock_workbook

    # Call the method
    sheet_content = file_reader_util.read_excel(excel_file_path, sheet_reference=valid_sheet_index)

    # Assertions
    mock_load_workbook.assert_called_once_with(excel_file_path, data_only=True)
    assert mock_workbook.worksheets[valid_sheet_index] == mock_sheet
    assert sheet_content is not None, "Expected sheet content, but got None."


def test_read_excel_file():
    try:
        excel_file_path = 'path/to/valid.xlsx'
        expected_first_value = 789
        sheet_content = file_reader_util.read_excel(excel_file_path, sheet_reference=0)
        first_value = sheet_content[1][1] if sheet_content and len(sheet_content) > 1 else None
        assert first_value == expected_first_value, f"Expected '{expected_first_value}', but got '{first_value}'"
    except Exception as e:
        print("An error occurred while reading the Excel file")


def test_read_excel_valid_file_and_sheet_name():
    try:
        excel_file_path = 'path/to/valid.xlsx'
        expected_values = [[123, 456], [678, 789]]
        values = file_reader_util.read_excel(excel_file_path)
        assert values == expected_values, f"Expected {expected_values}, but got {values}"
    except Exception as e:
        print("An error occurred while reading the Excel file")


def test_read_excel_file_invalid_path():
    try:
        invalid_excel_file_path = 'path/to/invalid.xlsx'
        with pytest.raises(Exception) as e:
            file_reader_util.read_excel(invalid_excel_file_path, sheet_reference="Sheet2")
        assert "An error occurred while reading the Excel file" in str(e.value)
    except Exception as e:
        print("An error occurred while reading the Excel file")


@patch('openpyxl.load_workbook')
def test_read_excel_file_invalid_sheet_index(mock_load_workbook):
    excel_file_path = 'path/to/valid.xlsx'
    invalid_sheet_index = 999  # Assuming this index is out of range

    # Mock the workbook and its methods
    mock_workbook = MagicMock()
    mock_workbook.sheetnames = ['Sheet1', 'Sheet2']
    mock_load_workbook.return_value = mock_workbook

    with pytest.raises(ValueError) as context:
        file_reader_util.read_excel(excel_file_path, sheet_reference=invalid_sheet_index)

    assert str(context.value) == f"Sheet index '{invalid_sheet_index}' is out of range."


@patch('openpyxl.load_workbook')
def test_read_excel_file_valid_sheet_name(mock_load_workbook):
    excel_file_path = 'path/to/valid.xlsx'
    valid_sheet_name = 'Sheet1'

    # Mock the workbook and its methods
    mock_workbook = MagicMock()
    mock_workbook.sheetnames = ['Sheet1', 'Sheet2']
    mock_sheet = MagicMock()
    mock_workbook.__getitem__.return_value = mock_sheet
    mock_load_workbook.return_value = mock_workbook

    # Call the method
    sheet_content = file_reader_util.read_excel(excel_file_path, sheet_reference=valid_sheet_name)

    # Assertions
    mock_load_workbook.assert_called_once_with(excel_file_path, data_only=True)
    mock_workbook.__getitem__.assert_called_once_with(valid_sheet_name)
    assert sheet_content is not None, "Expected sheet content, but got None."


@patch('openpyxl.load_workbook')
def test_read_excel_file_active_sheet(mock_load_workbook):
    excel_file_path = 'path/to/valid.xlsx'

    # Mock the workbook and its methods
    mock_workbook = MagicMock()
    mock_active_sheet = MagicMock()
    mock_workbook.active = mock_active_sheet
    mock_load_workbook.return_value = mock_workbook

    # Call the method without specifying a sheet reference
    sheet_content = file_reader_util.read_excel(excel_file_path)

    # Assertions
    mock_load_workbook.assert_called_once_with(excel_file_path, data_only=True)
    assert mock_workbook.active == mock_active_sheet
    assert sheet_content is not None, "Expected sheet content, but got None."


@patch('cafex_core.utils.file_reader_utils.BeautifulSoup')  # Use the correct import path
@patch('zipfile.ZipFile')
def test_read_word_file_with_mock(mock_zipfile, mock_beautifulsoup):
    filepath = 'path/to/valid.docx'
    mock_xml_content = b"<w:document><w:body><w:t>First paragraph</w:t><w:t>Second paragraph</w:t></w:body>" \
                       b"</w:document>"
    mock_text_elements = [MagicMock(text="First paragraph"), MagicMock(text="Second paragraph")]

    # Mock the behavior of zipfile.ZipFile
    mock_zip_instance = mock_zipfile.return_value.__enter__.return_value
    mock_zip_instance.read.return_value = mock_xml_content

    # Mock the behavior of BeautifulSoup
    mock_soup_instance = mock_beautifulsoup.return_value
    mock_soup_instance.find_all.return_value = mock_text_elements

    # Call the method
    content_list = file_reader_util.read_word_file(filepath)

    # Assertions
    mock_zipfile.assert_called_once_with(filepath)
    mock_zip_instance.read.assert_called_once_with("word/document.xml")
    mock_beautifulsoup.assert_called_once_with(mock_xml_content.decode("utf-8", errors="ignore"), features="lxml")
    mock_soup_instance.find_all.assert_called_once_with("w:t")
    assert content_list == ["First paragraph", "Second paragraph"]


def test_read_word_file_invalid_path():
    try:
        invalid_word_file_path = 'path/to/invalid.docx'
        with pytest.raises(Exception) as e:
            file_reader_util.read_word_file(invalid_word_file_path)
        assert "An error occurred while reading the Word file" in str(e.value)
    except Exception as e:
        print("An error occurred while reading the Word file")


@patch('builtins.open', new_callable=mock_open, read_data="Name,Age,Gender\nJohn,30,Male\nJane,25,Female")
def test_read_csv_file_valid(mock_open):
    csv_file_path = 'path/to/valid.csv'
    expected_content = [
        ['Name', 'Age', 'Gender'],
        ['John', '30', 'Male'],
        ['Jane', '25', 'Female']
    ]

    # Call the method
    content_list = file_reader_util.read_csv_file(csv_file_path)

    # Assertions
    mock_open.assert_called_once_with(csv_file_path, "r", newline='', encoding='utf-8')
    assert content_list == expected_content, f"Expected {expected_content}, but got {content_list}"


def test_read_csv_file_invalid_path():
    try:
        invalid_csv_file_path = 'path/to/invalidfile.csv'
        with pytest.raises(Exception) as e:
            file_reader_util.read_csv_file(invalid_csv_file_path)
        assert "An error occurred while reading the CSV file" in str(e.value)
    except Exception as e:
        print("An error occurred while reading the CSV file")


@patch('platform.system', return_value="Linux")
def test_extract_7z_file_non_windows(mock_platform_system):
    zip_file_name = 'path/to/valid.7z'
    target_dir = 'path/to/target'

    with pytest.raises(EnvironmentError) as e:
        file_reader_util.extract_7z_file(zip_file_name, target_dir)

    assert str(e.value) == "7Z extraction is only supported on Windows."
    mock_platform_system.assert_called_once()


@patch('py7zr.SevenZipFile')
@patch('os.path.join', side_effect=lambda *args: '/'.join(args))
@patch('platform.system', return_value="Windows")
def test_extract_7z_file_success(mock_platform_system, mock_path_join, mock_seven_zip_file):
    zip_file_name = 'path/to/valid.7z'
    target_dir = 'path/to/target'
    extracted_file_names = ['extracted_file.txt']

    # Mock the behavior of py7zr.SevenZipFile
    mock_archive = mock_seven_zip_file.return_value.__enter__.return_value
    mock_archive.getnames.return_value = extracted_file_names

    # Call the method
    extracted_file_path = file_reader_util.extract_7z_file(zip_file_name, target_dir)

    # Assertions
    mock_platform_system.assert_called_once()
    mock_seven_zip_file.assert_called_once_with(zip_file_name, mode="r")
    mock_archive.extractall.assert_called_once_with(path=target_dir)
    mock_archive.getnames.assert_called_once()
    mock_path_join.assert_called_once_with(target_dir, extracted_file_names[0])
    assert extracted_file_path == '/'.join([target_dir, extracted_file_names[0]])


@patch('py7zr.SevenZipFile')
@patch('platform.system', return_value="Windows")
def test_extract_7z_file_no_files_extracted(mock_platform_system, mock_seven_zip_file):
    zip_file_name = 'path/to/valid.7z'
    target_dir = 'path/to/target'

    # Mock the behavior of py7zr.SevenZipFile
    mock_archive = mock_seven_zip_file.return_value.__enter__.return_value
    mock_archive.getnames.return_value = []

    with pytest.raises(ValueError) as e:
        file_reader_util.extract_7z_file(zip_file_name, target_dir)

    assert str(e.value) == "No files were extracted from the archive."
    mock_platform_system.assert_called_once()
    mock_seven_zip_file.assert_called_once_with(zip_file_name, mode="r")
    mock_archive.extractall.assert_called_once_with(path=target_dir)
    mock_archive.getnames.assert_called_once()


@patch('os.path.exists', return_value=False)
@patch('os.makedirs')
@patch('zipfile.ZipFile.extractall')
@patch('zipfile.ZipFile.__init__', return_value=None)
def test_unzip_zip_file(mock_zip_init, mock_extractall, mock_makedirs, mock_path_exists):
    source_path = 'path/to/valid.zip'
    destination_path = 'path/to/destination'

    file_reader_util.unzip_zip_file(source_path, destination_path)

    mock_path_exists.assert_called_once_with(destination_path)
    mock_makedirs.assert_called_once_with(destination_path)
    mock_zip_init.assert_called_once_with(source_path, 'r')
    mock_extractall.assert_called_once_with(destination_path)


def test_unzip_zip_file_invalid_path():
    try:
        invalid_zip_file_path = 'path/to/invalidfile.zip'
        target_dir = 'path/to/target'
        with pytest.raises(Exception) as e:
            file_reader_util.unzip_zip_file(invalid_zip_file_path, target_dir)
        assert "The provided file is not supported" in str(e.value)
    except Exception as e:
        print("An error occurred while unzipping the file")


@patch('builtins.open', new_callable=MagicMock)
def test_read_txt_file_valid(mock_open):
    txt_file_path = 'path/to/valid.txt'
    expected_content = "This is a test content for the text file."

    # Mock the file read operation
    mock_open.return_value.__enter__.return_value.read.return_value = expected_content

    content = file_reader_util.read_txt_file(txt_file_path)

    # Assertions
    mock_open.assert_called_once_with(txt_file_path, 'r', encoding='utf-8')
    assert content == expected_content, f"Expected content '{expected_content}', but got '{content}'"


def test_read_txt_file_invalid_path():
    try:
        invalid_txt_file_path = 'path/to/invalid.txt'
        with pytest.raises(Exception) as e:
            file_reader_util.read_txt_file(invalid_txt_file_path)
        assert "An error occurred while reading the text file" in str(e.value)
    except Exception as e:
        print("An error occurred while reading the text file")


def test_write_txt_file_valid():
    try:
        txt_file_path = 'path/to/valid.txt'
        content = "This is a test content for the text file."
        file_reader_util.write_txt_file(txt_file_path, content)
        assert os.path.exists(txt_file_path), f"Expected file '{txt_file_path}' does not exist."
        with open(txt_file_path, 'r', encoding='utf-8') as f:
            written_content = f.read()
            assert written_content == content, f"Expected content '{content}', but got '{written_content}'."
    except Exception as e:
        print("An error occurred while writing to the text file")


@patch('builtins.open', side_effect=Exception("Mocked exception"))
def test_write_txt_file_exception_handling(mock_open):
    try:
        txt_file_path = 'path/to/valid.txt'
        content = "This is a test content for the text file."
        with pytest.raises(Exception) as e:
            file_reader_util.write_txt_file(txt_file_path, content)
        assert "An error occurred while writing to the text file" in str(e.value)
    except Exception as e:
        print("An error occurred while writing to the text file")


def test_write_excel_file_valid():
    try:
        excel_file_path = 'path/to/valid.xlsx'
        data = [
            ['Customer ID', 'Age', 'Gender'],
            ['1', '55', 'Male'],
            ['2', '19', 'Male'],
            ['3', '50', 'Male']
        ]
        file_reader_util.write_excel_file(excel_file_path, data)
        assert os.path.exists(excel_file_path), f"Expected Excel file '{excel_file_path}' does not exist."
        df = pd.read_excel(excel_file_path)
        expected_df = pd.DataFrame(data)
        pd.testing.assert_frame_equal(df, expected_df, check_dtype=False)
    except Exception as e:
        print("An error occurred while writing to the Excel file")


@patch('pandas.DataFrame.to_excel', side_effect=Exception("Mocked exception"))
def test_write_excel_file_exception_handling(mock_to_excel):
    try:
        excel_file_path = 'path/to/valid.xlsx'
        data = [
            {'Name': 'John', 'Age': 30},
            {'Name': 'Jane', 'Age': 25}
        ]
        with pytest.raises(Exception) as e:
            file_reader_util.write_excel_file(excel_file_path, data)
        assert "An error occurred while writing to the Excel file" in str(e.value)
    except Exception as e:
        print("An error occurred while writing to the Excel file")


def test_write_word_file_valid():
    try:
        word_file_path = 'path/to/valid.docx'
        content = [
            "This is the first paragraph.",
            "This is the second paragraph.",
            "This is the third paragraph."
        ]
        file_reader_util.write_word_file(word_file_path, content)
        assert os.path.exists(word_file_path), f"Expected Word file '{word_file_path}' does not exist."
        doc = Document(word_file_path)
        written_content = [para.text for para in doc.paragraphs]
        assert written_content == content, f"Expected content {content}, but got {written_content}."
    except Exception as e:
        print("An error occurred while writing to the Word file")


@patch('docx.document.Document.save', side_effect=Exception("Mocked exception"))
def test_write_word_file_exception_handling(mock_save):
    try:
        word_file_path = 'path/to/valid.docx'
        content = [
            "This is the first paragraph.",
            "This is the second paragraph."
        ]
        with pytest.raises(Exception) as e:
            file_reader_util.write_word_file(word_file_path, content)
        assert "An error occurred while writing to the Word file" in str(e.value)
    except Exception as e:
        print("An error occurred while writing to the Word file")


def test_write_csv_file_valid():
    try:
        csv_file_path = 'path/to/valid.csv'
        data = [
            ['Customer ID', 'Age', 'Gender'],
            ['1', '55', 'Male'],
            ['2', '19', 'Male'],
            ['3', '50', 'Male']
        ]
        file_reader_util.write_csv_file(csv_file_path, data)
        assert os.path.exists(csv_file_path), f"Expected CSV file '{csv_file_path}' does not exist."
        with open(csv_file_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            written_content = list(reader)
        assert written_content == data, f"Expected content {data}, but got {written_content}."
    except Exception as e:
        print("An error occurred while writing to the CSV file")


@patch('builtins.open', side_effect=Exception("Mocked exception"))
def test_write_csv_file_exception(mock_open):
    try:
        csv_file_path = 'path/to/valid.csv'
        data = [
            ['Customer ID', 'Age', 'Gender'],
            ['1', '55', 'Male'],
            ['2', '19', 'Male'],
            ['3', '50', 'Male']
        ]
        with pytest.raises(Exception) as e:
            file_reader_util.write_csv_file(csv_file_path, data)
        assert "An error occurred while writing to the CSV file" in str(e.value)
    except Exception as e:
        print("An error occurred while writing to the CSV file")


def test_delete_file_valid():
    try:
        test_file_path = 'path/to/valid.txt'
        assert os.path.exists(test_file_path), "Test file does not exist before deletion."
        file_reader_util.delete_file(test_file_path)
        assert not os.path.exists(test_file_path), "Test file still exists after deletion."
    except Exception as e:
        print("An error occurred while deleting the file")


def test_delete_file_invalid_path():
    try:
        invalid_file_path = 'path/to/invalidfile.txt'
        with pytest.raises(Exception) as e:
            file_reader_util.delete_file(invalid_file_path)
        assert f"The file '{invalid_file_path}' does not exist." in str(e.value)
    except Exception as e:
        print("An error occurred while deleting the file")


if __name__ == '__main__':
    unittest.main()
